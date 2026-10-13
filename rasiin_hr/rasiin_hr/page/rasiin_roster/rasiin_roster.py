import calendar
import datetime
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import add_days, cint, getdate, strip_html

DEFAULT_PAGE_LENGTH = 100

# Must match getRosterColor() in rasiin_roster.js  (background, border) without '#'
ROSTER_COLORS = {
	"Blue": ("EEF5FF", "9EC9FF"),
	"Cyan": ("ECFEFF", "67E8F9"),
	"Fuchsia": ("FDF4FF", "F0ABFC"),
	"Green": ("ECFDF5", "86EFAC"),
	"Lime": ("F7FEE7", "BEF264"),
	"Orange": ("FFF7ED", "FDBA74"),
	"Pink": ("FDF2F8", "F9A8D4"),
	"Red": ("FEF2F2", "FCA5A5"),
	"Violet": ("F5F3FF", "C4B5FD"),
	"Yellow": ("FEFCE8", "FDE047"),
}

# Cell values in the Excel grid that mean "no shift on this day"
OFF_VALUES = {"", "-", "off", "none", "x"}


# --------------------------------------------------------------------------- helpers


def _parse(value):
	if isinstance(value, str):
		value = frappe.parse_json(value)
	return frappe._dict(value or {})


def _format_time(value):
	"""Time fields come back as timedelta from MariaDB -> 'HH:MM'."""
	if not value:
		return ""
	if isinstance(value, datetime.timedelta):
		total = int(value.total_seconds())
		return f"{(total // 3600) % 24:02d}:{(total % 3600) // 60:02d}"
	return str(value)[:5]


def _month_range(filters):
	month, year = cint(filters.month), cint(filters.year)
	if not 1 <= month <= 12 or year < 1900:
		frappe.throw(_("Invalid month / year"))
	return datetime.date(year, month, 1), datetime.date(year, month, calendar.monthrange(year, month)[1])


def _daterange(start, end):
	d = start
	while d <= end:
		yield d
		d = d + datetime.timedelta(days=1)


def _extra_id_fields():
	meta = frappe.get_meta("Employee")
	return [f for f in ("employee_number", "attendance_device_id") if meta.has_field(f)]


def _get_employees(filters):
	"""All active employees matching the page filters + search (respects user permissions)."""
	extra = _extra_id_fields()

	emp_filters = [["status", "=", "Active"], ["company", "=", filters.company]]
	for field in ("department", "branch", "designation"):
		if filters.get(field):
			emp_filters.append([field, "=", filters.get(field)])

	or_filters = None
	search = (filters.get("employee_search") or "").strip()
	if search:
		like = f"%{search}%"
		or_filters = [[f, "like", like] for f in ["name", "employee_name", *extra]]

	return frappe.get_list(
		"Employee",
		filters=emp_filters,
		or_filters=or_filters,
		fields=["name", "employee_name", "department", "branch", "designation", "holiday_list", *extra],
		order_by="employee_name asc",
		limit_page_length=0,
	)


def _get_assignments(employee_names, start_date, end_date, shift_type=None):
	"""Submitted, Active assignments overlapping [start_date, end_date]."""
	if not employee_names:
		return []

	sa_filters = [
		["employee", "in", employee_names],
		["docstatus", "=", 1],
		["status", "=", "Active"],
		["start_date", "<=", end_date],
	]
	if shift_type:
		sa_filters.append(["shift_type", "=", shift_type])

	return frappe.get_list(
		"Shift Assignment",
		filters=sa_filters,
		or_filters=[["end_date", ">=", start_date], ["end_date", "is", "not set"]],
		fields=["name", "employee", "company", "shift_type", "start_date", "end_date", "status"],
		order_by="start_date asc",
		limit_page_length=0,
	)


def _get_shift_types():
	fields = ["name", "start_time", "end_time"]
	if frappe.get_meta("Shift Type").has_field("roster_color"):
		fields.append("roster_color")
	return frappe.get_all("Shift Type", fields=fields, order_by="name asc")


# --------------------------------------------------------------------------- roster data


@frappe.whitelist()
def get_roster_data(filters=None):
	filters = _parse(filters)

	if not filters.company:
		frappe.throw(_("Company is required"))

	start_date, end_date = _month_range(filters)
	limit = cint(filters.limit) or DEFAULT_PAGE_LENGTH

	all_employees = _get_employees(filters)
	total_employees = len(all_employees)
	employees = all_employees[:limit]

	assignments = _get_assignments([e.name for e in employees], start_date, end_date, filters.get("shift_type"))

	shift_map = {s.name: s for s in _get_shift_types()}
	for a in assignments:
		shift = shift_map.get(a.shift_type)
		a.start_date = str(a.start_date)
		a.end_date = str(a.end_date) if a.end_date else None
		a.start_time = _format_time(shift.start_time) if shift else ""
		a.end_time = _format_time(shift.end_time) if shift else ""
		a.roster_color = (shift.get("roster_color") if shift else "") or ""

	holidays = _get_holidays(employees, filters.company, start_date, end_date)

	dates = [
		{
			"date": str(d),
			"day": _(d.strftime("%a")),
			"label": d.strftime("%d"),
			"is_weekend": d.weekday() >= 5,
		}
		for d in _daterange(start_date, end_date)
	]

	return {
		"employees": employees,
		"total_employees": total_employees,
		"dates": dates,
		"assignments": assignments,
		"holidays": holidays,
	}


def _get_holidays(employees, company, start_date, end_date):
	default_list = frappe.get_cached_value("Company", company, "default_holiday_list")
	for e in employees:
		e.holiday_list = e.holiday_list or default_list

	lists = list({e.holiday_list for e in employees if e.holiday_list})
	if not lists:
		return {}

	holidays = {}
	for h in frappe.get_all(
		"Holiday",
		filters={"parent": ["in", lists], "holiday_date": ["between", [start_date, end_date]]},
		fields=["parent", "holiday_date", "description", "weekly_off"],
	):
		holidays.setdefault(h.parent, {})[str(h.holiday_date)] = {
			"description": strip_html(h.description or "").strip(),
			"weekly_off": cint(h.weekly_off),
		}
	return holidays


# --------------------------------------------------------------------------- excel: download grid


@frappe.whitelist()
def download_roster_template(filters=None):
	"""Excel grid like the roster page: employees down, days across, one shift per cell.
	Pre-filled with the current shifts so the department can edit and re-import."""
	from openpyxl import Workbook
	from openpyxl.formatting.rule import CellIsRule
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.datavalidation import DataValidation

	filters = _parse(filters)
	if not filters.company:
		frappe.throw(_("Company is required"))

	start_date, end_date = _month_range(filters)
	employees = _get_employees(filters)
	if not employees:
		frappe.throw(_("No employees found for the selected filters"))

	dates = list(_daterange(start_date, end_date))
	shift_types = _get_shift_types()
	holidays = _get_holidays(employees, filters.company, start_date, end_date)

	current = {}
	for a in _get_assignments([e.name for e in employees], start_date, end_date):
		a_end = getdate(a.end_date) if a.end_date else end_date
		for d in _daterange(max(getdate(a.start_date), start_date), min(a_end, end_date)):
			current.setdefault((a.employee, d), a.shift_type)

	wb = Workbook()
	ws = wb.active
	ws.title = "Roster"

	thin = Side(style="thin", color="D9D9D9")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	head_fill = PatternFill("solid", fgColor="F2F2F2")
	off_fill = PatternFill("solid", fgColor="FDE2E2")
	first_day_col, last_col = 3, 2 + len(dates)
	last_row = 2 + len(employees)

	ws.cell(1, 1, f"Roster – {filters.company} – {start_date.strftime('%B %Y')}").font = Font(bold=True, size=13)
	ws.cell(1, first_day_col, _("Pick a shift in each cell (or leave empty = no shift). Do not change the Employee ID column or the date row.")).font = Font(italic=True, color="808080")

	for col, label in ((1, "Employee ID"), (2, "Employee Name")):
		c = ws.cell(2, col, label)
		c.font, c.fill, c.border = Font(bold=True), head_fill, border
	for i, d in enumerate(dates):
		c = ws.cell(2, first_day_col + i, d)
		c.number_format = "ddd dd"
		c.font, c.fill, c.border = Font(bold=True), head_fill, border
		c.alignment = Alignment(horizontal="center")
		ws.column_dimensions[get_column_letter(first_day_col + i)].width = 11

	for r, emp in enumerate(employees, start=3):
		ws.cell(r, 1, emp.name).border = border
		ws.cell(r, 2, emp.employee_name).border = border
		emp_holidays = holidays.get(emp.holiday_list, {})
		for i, d in enumerate(dates):
			c = ws.cell(r, first_day_col + i, current.get((emp.name, d)))
			c.border = border
			c.alignment = Alignment(horizontal="center")
			if str(d) in emp_holidays:
				c.fill = off_fill

	ws.column_dimensions["A"].width = 17
	ws.column_dimensions["B"].width = 26
	ws.freeze_panes = ws.cell(3, first_day_col)

	# print nicely: landscape, all days on one page width, header rows repeated
	ws.page_setup.orientation = "landscape"
	ws.page_setup.paperSize = ws.PAPERSIZE_A3
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.print_title_rows = "1:2"

	grid = f"{get_column_letter(first_day_col)}3:{get_column_letter(last_col)}{last_row}"

	# Shift list sheet (dropdown source + legend)
	st = wb.create_sheet("Shift Types")
	st.append(["Shift Type", "Start", "End"])
	for c in st[1]:
		c.font = Font(bold=True)
	for s in shift_types:
		st.append([s.name, _format_time(s.start_time), _format_time(s.end_time)])
		color = ROSTER_COLORS.get(s.get("roster_color") or "Blue", ROSTER_COLORS["Blue"])
		st.cell(st.max_row, 1).fill = PatternFill("solid", fgColor=color[0])
		# colour the grid cells like the roster cards
		ws.conditional_formatting.add(
			grid,
			CellIsRule(
				operator="equal",
				formula=['"{}"'.format(s.name.replace('"', '""'))],
				fill=PatternFill(start_color=color[0], end_color=color[0], fill_type="solid"),
				font=Font(bold=True, color="111827"),
				border=Border(left=Side(style="thin", color=color[1]), right=Side(style="thin", color=color[1]),
					top=Side(style="thin", color=color[1]), bottom=Side(style="thin", color=color[1])),
			),
		)
	st.column_dimensions["A"].width = 24
	st.append([])
	st.append([_("Red-shaded cells on the Roster sheet are weekly offs / holidays.")])

	if shift_types:
		dv = DataValidation(
			type="list",
			formula1=f"='Shift Types'!$A$2:$A${len(shift_types) + 1}",
			allow_blank=True,
			showErrorMessage=True,
			errorTitle=_("Unknown shift"),
			error=_("Pick a shift from the list, or leave the cell empty."),
		)
		ws.add_data_validation(dv)
		dv.add(grid)

	buf = BytesIO()
	wb.save(buf)

	frappe.response["filename"] = f"Roster {filters.company} {start_date.strftime('%Y-%m')}.xlsx"
	frappe.response["filecontent"] = buf.getvalue()
	frappe.response["type"] = "binary"


# --------------------------------------------------------------------------- excel: import grid


@frappe.whitelist()
def import_roster(file_url, company, mode="Replace", dry_run=1):
	"""Import the Excel grid.

	mode = "Replace": for every employee in the file, the file IS the roster for the dates in it
	                  (empty cell = no shift). Existing shifts that already match are kept untouched.
	mode = "Add Only": only fill days that have no shift yet; existing shifts are never changed.
	dry_run = 1 returns a preview and changes nothing.
	"""
	dry_run = cint(dry_run)
	if mode not in ("Replace", "Add Only"):
		frappe.throw(_("Invalid mode"))
	if not company:
		frappe.throw(_("Company is required"))
	for ptype in ("create", "submit"):
		if not frappe.has_permission("Shift Assignment", ptype):
			frappe.throw(_("You need permission to {0} Shift Assignments").format(ptype), frappe.PermissionError)

	start_date, end_date, desired, errors = _read_roster_file(file_url, company)
	if errors:
		return {"errors": errors[:100], "error_count": len(errors)}

	plan = _plan_import(desired, start_date, end_date, mode)

	# Changing/removing existing shifts means cancelling them: only check that when it's needed
	if plan.cancel and not frappe.has_permission("Shift Assignment", "cancel"):
		return {
			"errors": [
				_("This file changes or removes {0} existing shift(s), which needs Cancel permission on Shift Assignment. "
				  "Use 'Add Only' mode, or ask HR to grant Cancel permission.").format(len(plan.cancel))
			],
			"error_count": 1,
		}
	summary = {
		"start_date": str(start_date),
		"end_date": str(end_date),
		"employees": len(desired),
		"create": len(plan.create),
		"cancel": len(plan.cancel),
		"keep": plan.keep,
		"skipped": plan.skipped,
		"changes": plan.changes[:300],
		"dry_run": dry_run,
	}
	if dry_run:
		return summary

	_execute_plan(plan, company, start_date, end_date)
	return summary


def _to_date(value):
	if value in (None, ""):
		return None
	if isinstance(value, datetime.datetime):
		return value.date()
	if isinstance(value, datetime.date):
		return value
	try:
		return getdate(str(value).strip())
	except Exception:
		return None


def _read_roster_file(file_url, company):
	from openpyxl import load_workbook

	if not (file_url or "").lower().endswith(".xlsx"):
		frappe.throw(_("Please upload the roster as an .xlsx file (the downloaded template)"))

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	wb = load_workbook(BytesIO(file_doc.get_content()), data_only=True)
	ws = wb["Roster"] if "Roster" in wb.sheetnames else wb.worksheets[0]

	errors = []

	# ---- find header row ("Employee ID" in column A)
	header_row = None
	for r in range(1, 11):
		if str(ws.cell(r, 1).value or "").strip().lower() in ("employee id", "employee"):
			header_row = r
			break
	if not header_row:
		frappe.throw(_("Could not find the 'Employee ID' header. Please use the template downloaded from the Roster page."))

	# ---- dates across (from column C)
	dates, col = [], 3
	while True:
		raw = ws.cell(header_row, col).value
		if raw in (None, ""):
			break
		d = _to_date(raw)
		if not d:
			errors.append(_("Header cell {0}: '{1}' is not a date").format(ws.cell(header_row, col).coordinate, raw))
		dates.append(d)
		col += 1
	if not dates:
		frappe.throw(_("No dates found in the header row"))
	if errors:
		return None, None, {}, errors
	for i in range(1, len(dates)):
		if dates[i] != add_days(dates[0], i):
			frappe.throw(_("Dates in the header row must be consecutive days (problem at {0})").format(dates[i]))
	start_date, end_date = dates[0], dates[-1]

	# ---- lookups
	shift_lookup = {s.name.lower(): s.name for s in frappe.get_all("Shift Type", fields=["name"])}
	extra = _extra_id_fields()

	rows = []
	for r in range(header_row + 1, ws.max_row + 1):
		emp_id = str(ws.cell(r, 1).value or "").strip()
		cells = [ws.cell(r, 3 + i).value for i in range(len(dates))]
		if not emp_id and all(str(v or "").strip() == "" for v in cells):
			continue  # blank line
		rows.append((r, emp_id, cells))

	ids = [x[1] for x in rows if x[1]]
	emp_fields = ["name", "employee_name", "company", "status", *extra]
	found = frappe.get_all("Employee", filters={"name": ["in", ids]}, fields=emp_fields) if ids else []
	for f in extra:  # allow Employee Number / Device ID in the ID column too
		if ids:
			found += frappe.get_all("Employee", filters={f: ["in", ids]}, fields=emp_fields)
	by_key = {}
	for e in found:
		for key in [e.name, *[e.get(f) for f in extra]]:
			if key:
				by_key.setdefault(str(key), e)

	desired = {}
	for r, emp_id, cells in rows:
		if not emp_id:
			errors.append(_("Row {0}: Employee ID is empty").format(r))
			continue
		emp = by_key.get(emp_id)
		if not emp:
			errors.append(_("Row {0}: Employee '{1}' not found").format(r, emp_id))
			continue
		if emp.company != company:
			errors.append(_("Row {0}: {1} belongs to company {2}, not {3}").format(r, emp.name, emp.company, company))
			continue
		if emp.status != "Active":
			errors.append(_("Row {0}: {1} ({2}) is {3}").format(r, emp.name, emp.employee_name, emp.status))
			continue
		if not frappe.has_permission("Employee", "read", emp.name):
			errors.append(_("Row {0}: no permission for employee {1}").format(r, emp.name))
			continue
		if emp.name in desired:
			errors.append(_("Row {0}: {1} appears more than once in the file").format(r, emp.name))
			continue

		want = {}
		for i, value in enumerate(cells):
			text = str(value or "").strip()
			if text.lower() in OFF_VALUES:
				continue
			shift = shift_lookup.get(text.lower())
			if not shift:
				errors.append(
					_("Row {0} ({1}), {2}: unknown shift '{3}'").format(r, emp.employee_name, dates[i].strftime("%d %b"), text)
				)
				continue
			want[dates[i]] = shift
		desired[emp.name] = frappe._dict(employee_name=emp.employee_name, want=want)

	return start_date, end_date, desired, errors


def _runs(want):
	"""{date: shift} -> [(shift, start, end)] grouping consecutive days with the same shift."""
	runs = []
	for d in sorted(want):
		if runs and runs[-1][0] == want[d] and runs[-1][2] == d - datetime.timedelta(days=1):
			runs[-1][2] = d
		else:
			runs.append([want[d], d, d])
	return [tuple(r) for r in runs]


def _fmt(d):
	return d.strftime("%d %b") if d else _("open-ended")


def _plan_import(desired, start_date, end_date, mode):
	plan = frappe._dict(create=[], cancel=[], keep=0, skipped=0, changes=[])
	existing = {}
	for a in _get_assignments(list(desired), start_date, end_date):
		existing.setdefault(a.employee, []).append(a)

	for emp, info in desired.items():
		want = dict(info.want)
		label = f"{emp} {info.employee_name}"

		for a in existing.get(emp, []):
			a_start, a_end = getdate(a.start_date), getdate(a.end_date) if a.end_date else None
			days = list(_daterange(max(a_start, start_date), min(a_end or end_date, end_date)))

			if mode == "Add Only":
				for d in days:
					if d in want:
						if want[d] != a.shift_type:
							plan.skipped += 1
						want.pop(d)
				continue

			if all(want.get(d) == a.shift_type for d in days):
				plan.keep += 1  # already exactly right: leave it alone
				for d in days:
					want.pop(d, None)
			else:
				plan.cancel.append(a)
				plan.changes.append(
					f"− {label}: {a.shift_type} {_fmt(max(a_start, start_date))} → {_fmt(min(a_end or end_date, end_date))} ({a.name})"
				)

		for shift, s, e in _runs(want):
			plan.create.append(frappe._dict(employee=emp, shift_type=shift, start_date=s, end_date=e))
			plan.changes.append(f"+ {label}: {shift} {_fmt(s)} → {_fmt(e)}")

	return plan


def _execute_plan(plan, company, start_date, end_date):
	"""Runs inside the request transaction: any failure rolls everything back."""
	context = ""
	try:
		for a in plan.cancel:
			context = f"{a.employee} {a.name}"
			doc = frappe.get_doc("Shift Assignment", a.name)
			a_start = getdate(doc.start_date)
			a_end = getdate(doc.end_date) if doc.end_date else None
			doc.cancel()
			# keep the parts of the old assignment that fall outside the imported dates
			if a_start < start_date:
				create_replacement_shift(doc.employee, doc.company, doc.shift_type, a_start, add_days(start_date, -1), doc.status)
			if a_end is None or a_end > end_date:
				create_replacement_shift(doc.employee, doc.company, doc.shift_type, add_days(end_date, 1), a_end, doc.status)

		for c in plan.create:
			context = f"{c.employee} {c.shift_type} {c.start_date} → {c.end_date}"
			create_replacement_shift(c.employee, company, c.shift_type, c.start_date, c.end_date)
	except Exception as e:
		frappe.db.rollback()
		frappe.clear_messages()
		frappe.throw(
			_("Import stopped and nothing was saved. Problem at {0}: {1}").format(context, strip_html(str(e))),
			title=_("Roster Import Failed"),
		)


# --------------------------------------------------------------------------- create / remove


@frappe.whitelist()
def create_shift_assignment(args=None):
	args = _parse(args)

	for field in ("employee", "company", "shift_type", "start_date"):
		if not args.get(field):
			frappe.throw(_("{0} is required").format(frappe.unscrub(field)))

	start_date = getdate(args.start_date)
	end_date = getdate(args.end_date) if args.end_date else None

	if end_date and end_date < start_date:
		frappe.throw(_("End Date cannot be before Start Date"))

	conflict = get_conflicting_shift_assignment(args.employee, start_date, end_date, args.get("ignore_assignment"))
	if conflict:
		frappe.throw(
			_("Employee already has active shift {0} ({1}) overlapping this period.").format(
				frappe.bold(conflict.shift_type), frappe.utils.get_link_to_form("Shift Assignment", conflict.name)
			)
		)

	return create_replacement_shift(
		args.employee,
		args.company,
		args.shift_type,
		start_date,
		end_date,
		args.get("status") or "Active",
	)


@frappe.whitelist()
def cancel_shift_assignment(name):
	doc = frappe.get_doc("Shift Assignment", name)
	doc.cancel()
	return True


def get_conflicting_shift_assignment(employee, start_date, end_date=None, ignore_assignment=None):
	"""end_date=None means open-ended (overlaps everything after start_date)."""
	rows = frappe.db.sql(
		"""
		SELECT name, shift_type, start_date, end_date
		FROM `tabShift Assignment`
		WHERE employee = %(employee)s
		  AND docstatus = 1
		  AND status = 'Active'
		  AND name != %(ignore_assignment)s
		  AND start_date <= %(end_date)s
		  AND IFNULL(end_date, '9999-12-31') >= %(start_date)s
		LIMIT 1
		""",
		{
			"employee": employee,
			"start_date": getdate(start_date),
			"end_date": getdate(end_date) if end_date else getdate("9999-12-31"),
			"ignore_assignment": ignore_assignment or "",
		},
		as_dict=True,
	)
	return rows[0] if rows else None


@frappe.whitelist()
def remove_shift_assignment_range(assignment, from_date, to_date=None):
	"""Remove [from_date, to_date] from an assignment, keeping the parts before/after.
	to_date empty => remove until the end of the assignment (works for open-ended ones too)."""
	old = frappe.get_doc("Shift Assignment", assignment)

	if old.docstatus != 1:
		frappe.throw(_("Only a submitted Shift Assignment can be removed"))

	old_start = getdate(old.start_date)
	old_end = getdate(old.end_date) if old.end_date else None  # None = open-ended

	from_date = getdate(from_date)
	to_date = getdate(to_date) if to_date else old_end

	if to_date and to_date < from_date:
		frappe.throw(_("To Date cannot be before From Date"))

	if from_date < old_start or (old_end and to_date and to_date > old_end):
		frappe.throw(_("Selected date range is outside this shift assignment"))

	keep = frappe._dict(
		employee=old.employee,
		company=old.company,
		shift_type=old.shift_type,
		status=old.status,
	)

	# Whole request runs in one transaction: if a replacement fails, the cancel is rolled back too.
	old.cancel()

	if old_start < from_date:
		create_replacement_shift(keep.employee, keep.company, keep.shift_type, old_start, add_days(from_date, -1), keep.status)

	# something remains after the removed range (to_date=None means removed till the end)
	if to_date and (old_end is None or to_date < old_end):
		create_replacement_shift(keep.employee, keep.company, keep.shift_type, add_days(to_date, 1), old_end, keep.status)

	return True


def create_replacement_shift(employee, company, shift_type, start_date, end_date=None, status="Active"):
	doc = frappe.new_doc("Shift Assignment")
	doc.employee = employee
	doc.company = company
	doc.shift_type = shift_type
	doc.start_date = start_date
	doc.end_date = end_date
	doc.status = status or "Active"

	doc.insert()
	doc.submit()
	return doc.name